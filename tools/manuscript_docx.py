"""Build a manuscript .docx from chapter markdown files.

Times New Roman 12, left aligned, no indents, 0 pt before/after,
1.15 line spacing. Chapter headings are outline level 1. Scene breaks
("---") become a "#" line.
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

FONT = "Times New Roman"


def fmt(par):
    pf = par.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15


def style_run(run):
    run.font.name = FONT
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def add_inline(par, text):
    # **bold** and *italic*
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*"):
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            r = par.add_run(tok)
        style_run(r)


def set_outline_level(par, level):
    ppr = par._element.get_or_add_pPr()
    ol = OxmlElement("w:outlineLvl")
    ol.set(qn("w:val"), str(level))
    ppr.append(ol)


def main(out, paths):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    first = True
    for path in paths:
        text = open(path, encoding="utf-8").read()
        blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
        for b in blocks:
            b = " ".join(line.strip() for line in b.splitlines())
            if b.startswith("# "):
                if not first:
                    doc.add_page_break()
                first = False
                p = doc.add_paragraph()
                add_inline(p, b[2:].strip())
                set_outline_level(p, 0)
            elif re.fullmatch(r"-{3,}|#|\*\s*\*\s*\*", b):
                p = doc.add_paragraph()
                add_inline(p, "#")
            else:
                p = doc.add_paragraph()
                add_inline(p, b)
            fmt(p)
    doc.save(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2:])
